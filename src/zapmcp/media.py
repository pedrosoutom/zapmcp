"""Anexos: imagens e documentos lidos direto da pasta do WhatsApp, sob demanda.

Diferente do banco, um arquivo comum pode ser lido sem cópia: não há travas a disputar com o
WhatsApp. Nada é escrito lá, e nada é guardado em disco do nosso lado.
"""

import io
import zipfile
from collections.abc import Iterable, Iterator
from dataclasses import dataclass
from functools import lru_cache
from html.parser import HTMLParser
from pathlib import Path

from fastmcp.utilities.types import Image

from . import config

# Anexos que o Claude consegue abrir pelas tools. Vídeo e figurinha ficam só no marcador.
OPENABLE = ("imagem", "documento")

# Acima disso o Claude reduz a imagem de qualquer jeito; mandar maior só gasta token.
IMAGE_MAX_SIDE = 1568
# Quanto texto vai em cada parte de read_document.
PART_CHARS = 15_000
# Páginas escaneadas (PDF sem texto) viram imagem; no máximo estas por parte.
MAX_SCANNED_PER_PART = 3
# Acima disso o documento não é aberto. ZIP só é listado, então não entra no limite.
MAX_DOCUMENT_BYTES = 60 * 1024 * 1024
MAX_SHEET_ROWS = 2_000
MAX_ZIP_ENTRIES = 500

IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp", ".gif"}
TEXT_SUFFIXES = {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".yaml", ".yml", ".log", ".ofx", ".vcf", ".mobileconfig"}
SUPPORTED = "PDF, Word (.docx), Excel (.xlsx), PowerPoint (.pptx), HTML, texto e ZIP"


class UnsupportedDocument(ValueError):
    pass


@dataclass(frozen=True)
class Part:
    text: str
    scanned_pages: tuple[int, ...] = ()


@dataclass(frozen=True)
class Document:
    kind: str
    extent: str
    parts: tuple[Part, ...]


def local_file(media_path: str | None) -> Path | None:
    """Caminho do anexo, se o arquivo existe e está dentro da pasta de mídias do WhatsApp."""
    if not media_path:
        return None
    try:
        root = (config.wa_dir() / "Message").resolve()
        path = (root / media_path).resolve()
        return path if path.is_relative_to(root) and path.is_file() else None
    except OSError:
        return None


def human_size(n: int) -> str:
    if n < 1024 * 1024:
        return f"{max(1, round(n / 1024))} KB"
    return f"{n / 1024 / 1024:.1f} MB".replace(".", ",")


# --- imagens ---


def image(path: Path) -> Image:
    from PIL import Image as PILImage

    with PILImage.open(path) as img:
        return _to_image(img)


def pdf_page_image(path: Path, index: int) -> Image:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(path)
    try:
        page = pdf[index]
        # Tamanho da página em pontos (1/72 pol.): escala para o lado maior ficar em IMAGE_MAX_SIDE.
        scale = min(IMAGE_MAX_SIDE / max(page.get_size()), 4)
        return _to_image(page.render(scale=scale).to_pil())
    finally:
        pdf.close()


def _to_image(img) -> Image:
    from PIL import ImageOps

    img = ImageOps.exif_transpose(img)
    if max(img.size) > IMAGE_MAX_SIDE:
        img.thumbnail((IMAGE_MAX_SIDE, IMAGE_MAX_SIDE))
    if img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=85)
    return Image(data=buf.getvalue(), format="jpeg")


# --- documentos ---


def read_document(path: Path) -> Document:
    st = path.stat()
    return _read_cached(str(path), st.st_mtime_ns, st.st_size)


@lru_cache(maxsize=32)
def _read_cached(path_str: str, _mtime: int, size: int) -> Document:
    path = Path(path_str)
    suffix = path.suffix.lower()
    if suffix in (".zip", ".skill"):
        return _zip(path)
    readers = {".pdf": _pdf, ".docx": _docx, ".xlsx": _xlsx, ".pptx": _pptx, ".html": _html, ".htm": _html}
    if suffix not in readers and suffix not in TEXT_SUFFIXES:
        raise UnsupportedDocument(f"Formato {suffix or 'sem extensão'} não suportado. Suportados: {SUPPORTED}.")
    if size > MAX_DOCUMENT_BYTES:
        raise UnsupportedDocument(f"Documento grande demais para abrir ({human_size(size)}).")
    if suffix in readers:
        return readers[suffix](path)
    return Document("texto", human_size(size), _pack([(path.read_text(errors="replace"), None)]))


def _pdf(path: Path) -> Document:
    import pypdfium2 as pdfium

    try:
        pdf = pdfium.PdfDocument(path)
    except pdfium.PdfiumError as e:
        if "password" in str(e).lower():
            raise UnsupportedDocument("PDF protegido por senha.") from e
        raise
    try:
        pages = []
        for page in pdf:
            textpage = page.get_textpage()
            pages.append(textpage.get_text_range().strip())
            textpage.close()
            page.close()
    finally:
        pdf.close()
    # Página quase sem texto é, na prática, escaneada: vai como imagem.
    sections = [
        (f"--- página {i} ---\n{text}", None) if len(text) >= 20 else (f"--- página {i} (escaneada, vai como imagem) ---", i - 1)
        for i, text in enumerate(pages, 1)
    ]
    return Document("PDF", _count(len(pages), "página", "páginas"), _pack(sections))


def _docx(path: Path) -> Document:
    import docx

    doc = docx.Document(path)
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        lines.append("")
        lines += [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
    return Document("Word", _count(len(doc.paragraphs), "parágrafo", "parágrafos"), _pack([("\n".join(lines), None)]))


def _xlsx(path: Path) -> Document:
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sections = []
        for sheet in book.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                if len(rows) == MAX_SHEET_ROWS:
                    rows.append(f"… aba cortada em {MAX_SHEET_ROWS} linhas")
                    break
                if any(v is not None for v in row):
                    rows.append("\t".join("" if v is None else str(v) for v in row).rstrip())
            sections.append((f"## Aba: {sheet.title}\n" + "\n".join(rows), None))
    finally:
        book.close()
    return Document("Excel", _count(len(sections), "aba", "abas"), _pack(sections))


def _pptx(path: Path) -> Document:
    import pptx

    deck = pptx.Presentation(path)
    sections = []
    for i, slide in enumerate(deck.slides, 1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text)
            if shape.has_table:
                lines += [" | ".join(c.text.strip() for c in row.cells) for row in shape.table.rows]
        if slide.has_notes_slide and (notes := slide.notes_slide.notes_text_frame.text.strip()):
            lines.append(f"[notas] {notes}")
        sections.append((f"--- slide {i} ---\n" + "\n".join(lines), None))
    return Document("PowerPoint", _count(len(sections), "slide", "slides"), _pack(sections))


class _HTMLText(HTMLParser):
    BLOCKS = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "table"}

    def __init__(self) -> None:
        super().__init__()
        self.chunks: list[str] = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "noscript"):
            self.hidden += 1
        elif tag in self.BLOCKS:
            self.chunks.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript"):
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, data):
        if not self.hidden:
            self.chunks.append(data)


def _html(path: Path) -> Document:
    parser = _HTMLText()
    parser.feed(path.read_text(errors="replace"))
    lines = [" ".join(line.split()) for line in "".join(parser.chunks).splitlines()]
    text = "\n".join(line for line in lines if line)
    return Document("HTML", human_size(path.stat().st_size), _pack([(text, None)]))


def _zip(path: Path) -> Document:
    with zipfile.ZipFile(path) as archive:
        entries = [e for e in archive.infolist() if not e.is_dir()]
    lines = [f"{e.filename}  ({human_size(e.file_size)})" for e in entries[:MAX_ZIP_ENTRIES]]
    if len(entries) > MAX_ZIP_ENTRIES:
        lines.append(f"… e mais {len(entries) - MAX_ZIP_ENTRIES} arquivos")
    text = "Arquivo compactado. Conteúdo (os arquivos de dentro não são abertos):\n" + "\n".join(lines)
    return Document("ZIP", _count(len(entries), "arquivo", "arquivos"), _pack([(text, None)]))


def _pack(sections: Iterable[tuple[str, int | None]]) -> tuple[Part, ...]:
    """Junta seções (páginas, slides, abas) em partes de até PART_CHARS caracteres."""
    parts: list[Part] = []
    texts: list[str] = []
    scanned: list[int] = []
    size = 0
    for section, scanned_page in sections:
        for piece in _split(section, PART_CHARS):
            full = size + len(piece) > PART_CHARS
            too_many_images = scanned_page is not None and len(scanned) == MAX_SCANNED_PER_PART
            if (full or too_many_images) and texts:
                parts.append(Part("\n\n".join(texts), tuple(scanned)))
                texts, scanned, size = [], [], 0
            texts.append(piece)
            size += len(piece)
        if scanned_page is not None:
            scanned.append(scanned_page)
    if texts:
        parts.append(Part("\n\n".join(texts), tuple(scanned)))
    return tuple(parts) or (Part("(documento sem texto)"),)


def _count(n: int, one: str, many: str) -> str:
    return f"{n} {one if n == 1 else many}"


def _split(text: str, limit: int) -> Iterator[str]:
    """Quebra um texto longo em pedaços de até `limit`, de preferência no fim de uma linha."""
    while len(text) > limit:
        cut = text.rfind("\n", 0, limit)
        cut = cut if cut > limit // 2 else limit
        yield text[:cut]
        text = text[cut:].lstrip("\n")
    yield text
